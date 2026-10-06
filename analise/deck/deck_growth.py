"""Builder do deck padrão Growth/CRO - Grupo Oscar."""
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

TEXTO = RGBColor(0x59, 0x59, 0x59)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
CLARO = RGBColor(0xEF, 0xEF, 0xEF)
TRACO = RGBColor(0xEF, 0xEF, 0xEF)
FONTE_RODAPE = RGBColor(0xB7, 0xB7, 0xB7)
CINZA_PAINEL = RGBColor(0xF3, 0xF3, 0xF3)
AZUL_SUMARIO = RGBColor(0x42, 0x85, 0xF4)
BULLET = RGBColor(0x12, 0xD0, 0x00)
LINK = RGBColor(0x22, 0x00, 0xCC)

BOLD = "Montserrat"
BODY = "Montserrat Medium"
W, H = Inches(10), Inches(5.625)


def _in(v):
    return Emu(int(v * 914400))


def _sem_sombra(shape):
    from lxml import etree
    el = shape._element
    spPr = el.spPr
    for e in spPr.findall(qn('a:effectLst')):
        spPr.remove(e)
    etree.SubElement(spPr, qn('a:effectLst'))
    for e in el.findall(qn('p:style')):
        el.remove(e)


class Deck:
    def __init__(self, kicker="GROWTH | CRO", subtitulo_rodape=""):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.kicker = kicker
        self.subtitulo_rodape = subtitulo_rodape

    def _blank(self, bg=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        if bg is not None:
            self._set_bg(s, bg)
        return s

    @staticmethod
    def _set_bg(slide, rgb):
        from lxml import etree
        xml = ('<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
               'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
               f'<p:bgPr><a:solidFill><a:srgbClr val="{rgb}"/></a:solidFill>'
               '<a:effectLst/></p:bgPr></p:bg>')
        slide._element.find(qn('p:cSld')).insert(0, etree.fromstring(xml))

    def _tb(self, slide, x, y, w, h):
        tb = slide.shapes.add_textbox(_in(x), _in(y), _in(w), _in(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(91425)
        p0 = tf._txBody.find(qn('a:p'))
        p0.getparent().remove(p0)
        return tf

    @staticmethod
    def _p(tf, space_before=0, align=PP_ALIGN.LEFT):
        from lxml import etree
        from pptx.text.text import _Paragraph
        par = _Paragraph(etree.SubElement(tf._txBody, qn('a:p')), tf)
        par.alignment = align
        par.space_before = Pt(space_before)
        par.space_after = Pt(0)
        return par

    @staticmethod
    def _run(par, texto, size, bold=False, italic=False, color=TEXTO, font=None):
        r = par.add_run()
        r.text = texto
        f = r.font
        f.name = font or (BOLD if bold else BODY)
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = color
        rPr = r._r.get_or_add_rPr()
        for tag in ('a:ea', 'a:cs', 'a:sym'):
            rPr.append(rPr.makeelement(qn(tag), {'typeface': f.name}))
        return r

    def _bullet(self, par, char="●", size=10, color=BULLET, font=BODY):
        from lxml import etree
        pPr = par._p.get_or_add_pPr()
        pPr.set('marL', '457200')
        pPr.set('indent', '-292100')
        buClr = etree.SubElement(pPr, qn('a:buClr'))
        etree.SubElement(buClr, qn('a:srgbClr')).set('val', str(color))
        etree.SubElement(pPr, qn('a:buSzPts')).set('val', str(int(size * 100)))
        etree.SubElement(pPr, qn('a:buFont')).set('typeface', font)
        etree.SubElement(pPr, qn('a:buChar')).set('char', char)

    def _numero(self, par, size=14, color=TEXTO, font=BODY):
        from lxml import etree
        pPr = par._p.get_or_add_pPr()
        pPr.set('marL', '457200')
        pPr.set('indent', '-317500')
        buClr = etree.SubElement(pPr, qn('a:buClr'))
        etree.SubElement(buClr, qn('a:srgbClr')).set('val', str(color))
        etree.SubElement(pPr, qn('a:buSzPts')).set('val', str(int(size * 100)))
        etree.SubElement(pPr, qn('a:buFont')).set('typeface', font)
        etree.SubElement(pPr, qn('a:buAutoNum')).set('type', 'arabicPeriod')

    def _traco(self, slide, x, y, w):
        ln = slide.shapes.add_connector(1, _in(x), _in(y), _in(x + w), _in(y))
        ln.line.color.rgb = TRACO
        ln.line.width = Pt(1.5)
        _sem_sombra(ln)
        return ln

    @staticmethod
    def _fit(slide, path, cx, cy, cw, ch):
        """Imagem centralizada na area (cx,cy,cw,ch), sem distorcer."""
        from PIL import Image
        iw, ih = Image.open(path).size
        r = min(cw / iw, ch / ih)
        w, h = iw * r, ih * r
        return slide.shapes.add_picture(path, _in(cx + (cw - w) / 2),
                                        _in(cy + (ch - h) / 2), _in(w), _in(h))

    # ---------- layouts ----------
    def capa(self, titulo, subtitulo="", data="", imagem=None, logo=None):
        s = self._blank("B70701")
        if imagem:
            s.shapes.add_picture(imagem, _in(0), _in(-0.05), _in(10.13), _in(5.72))
        if logo:
            s.shapes.add_picture(logo, _in(0.48), _in(0.73), _in(1.03), _in(0.35))
        tf = self._tb(s, 0.36, 1.84, 7.24, 2.38)
        self._run(self._p(tf), titulo, 40, bold=True, color=BRANCO)
        if subtitulo:
            self._run(self._p(tf, 14), subtitulo, 18, color=BRANCO)
        tf = self._tb(s, 0.36, 4.64, 6.97, 0.82)
        self._run(self._p(tf), self.kicker, 18, color=CLARO)
        if data:
            self._run(self._p(tf), data, 18, color=CLARO)
        return s

    def sumario(self, itens, titulo="Sumário"):
        s = self._blank()
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, _in(-0.08), _in(-0.03),
                                 _in(2.84), _in(5.69))
        bar.fill.solid()
        bar.fill.fore_color.rgb = AZUL_SUMARIO
        bar.line.fill.background()
        _sem_sombra(bar)
        tf = self._tb(s, 3.33, 0.33, 6.29, 0.69)
        self._run(self._p(tf), titulo, 24, bold=True)
        self._traco(s, 3.41, 1.02, 6.13)
        tf = self._tb(s, 3.41, 1.37, 6.29, 2.25)
        for item in itens:
            par = self._p(tf, 10)
            self._numero(par)
            self._run(par, item, 14)
        return s

    def divisor(self, titulo, url=None, rodape=None):
        s = self._blank("434343")
        tf = self._tb(s, 0.36, 1.6, 5.21, 1.95)
        self._run(self._p(tf), titulo, 48, bold=True, color=BRANCO)
        if url:
            tf = self._tb(s, 0.46, 3.8, 5.21, 0.53)
            r = self._run(self._p(tf), url, 12, bold=True, color=LINK)
            r.hyperlink.address = url
        tf = self._tb(s, 0.36, 4.57, 6.97, 0.72)
        self._run(self._p(tf), self.kicker, 12, bold=True, color=BRANCO)
        self._run(self._p(tf), rodape or self.subtitulo_rodape, 12, bold=True, color=BRANCO)
        return s

    def _titulo_traco(self, s, titulo, x=0.33, y=0.33, w=9.25,
                      lx=0.45, ly=1.02, lw=9.02):
        tf = self._tb(s, x, y, w, 0.69)
        self._run(self._p(tf), titulo, 24, bold=True)
        self._traco(s, lx, ly, lw)

    def conteudo(self, titulo, corpo=None, bullets=None, imagem=None, fonte=None):
        s = self._blank()
        self._titulo_traco(s, titulo)
        if corpo or bullets:
            tf = self._tb(s, 0.45, 1.31, 9.25, 2.76)
            self._paragrafos(tf, corpo or [], bullets or [])
        if imagem:
            topo = 2.02 if (corpo or bullets) else 1.31
            base = 4.85 if fonte else 5.3
            self._fit(s, imagem, 0.45, topo, 9.1, base - topo)
        if fonte:
            tf = self._tb(s, 0.45, 4.98, 9.25, 0.34)
            self._run(self._p(tf), fonte, 10, color=FONTE_RODAPE)
        return s

    def split(self, titulo, corpo=None, bullets=None, imagem=None):
        s = self._blank()
        painel = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _in(5.605),
                                    _in(0.07), _in(5.24), _in(5.415))
        painel.adjustments[0] = 0.03975
        painel.fill.solid()
        painel.fill.fore_color.rgb = CINZA_PAINEL
        painel.line.fill.background()
        _sem_sombra(painel)
        self._titulo_traco(s, titulo, x=0.33, y=1.31, w=4.45, lx=0.39, ly=2.0, lw=4.34)
        if corpo or bullets:
            tf = self._tb(s, 0.39, 2.28, 4.45, 2.4)
            self._paragrafos(tf, corpo or [], bullets or [])
        if imagem:
            self._fit(s, imagem, 5.8, 0.4, 3.95, 4.75)
        return s

    def _paragrafos(self, tf, corpo, bullets):
        primeiro = True
        for bloco in corpo:
            par = self._p(tf, 10 if primeiro else 14)
            primeiro = False
            if isinstance(bloco, tuple):
                self._run(par, bloco[0], 13, bold=True)
                if len(bloco) > 1 and bloco[1]:
                    self._run(par, bloco[1], 13, bold=True, italic=True)
            else:
                self._run(par, bloco, 13)
        for b in bullets:
            par = self._p(tf, 14)
            self._bullet(par)
            self._run(par, b, 13)

    def save(self, path):
        self.prs.save(path)
        return path
